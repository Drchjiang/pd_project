# -*- coding: utf-8 -*-
from docx import Document
from docx.oxml.ns import qn

base = r"D:\内推\社招\chjiang\resume_template"
zh = Document(base + r"\数字后端工程师简历模板.docx")
en = Document(base + r"\数字后端工程师简历模板(英文版).docx")

zs = [p.style.name for p in zh.paragraphs]
es = [p.style.name for p in en.paragraphs]
print("中文段落:", len(zs), "表格:", len(zh.tables), "| 英文段落:", len(es), "表格:", len(en.tables))
print("段落样式序列完全一致(结构同步):", zs == es)

def titles(d):
    out = []
    for p in d.paragraphs:
        pPr = p._p.find(qn("w:pPr"))
        if pPr is not None and pPr.find(qn("w:pBdr")) is not None:
            out.append(p.text)
    return out
for a, b in zip(titles(zh), titles(en)):
    print("  章节:", a, " <-> ", b)

print("bullet 数 中/英:", zs.count("List Bullet"), "/", es.count("List Bullet"))

# 证件照框
for name, d in [("中文", zh), ("英文", en)]:
    h = d.tables[0]
    row = h.rows[0]
    trPr = row._tr.find(qn("w:trPr"))
    th = trPr.find(qn("w:trHeight"))
    lc, rc = row.cells
    tcPr = rc._tc.find(qn("w:tcPr"))
    tcB = tcPr.find(qn("w:tcBorders"))
    va = tcPr.find(qn("w:vAlign"))
    val = th.get(qn("w:val")); rule = th.get(qn("w:hRule"))
    print(f"[{name}] 行高twips={val} hRule={rule} 四边边框={tcB is not None} 垂直居中={va is not None} 照片格宽cm={round(rc.width.cm,2)}")
    print("   照片格文字:", " / ".join(p.text for p in rc.paragraphs))
    print("   姓名:", lc.paragraphs[0].text)

# 字体生效值：是否所有 run 都有 eastAsia + size
for name, d in [("中文", zh), ("英文", en)]:
    miss = 0
    for p in d.paragraphs:
        for r in p.runs:
            rPr = r._element.find(qn("w:rPr"))
            rf = rPr.find(qn("w:rFonts")) if rPr is not None else None
            if rf is None or rf.get(qn("w:eastAsia")) is None or r.font.size is None:
                miss += 1
    print(f"[{name}] 缺字体/字号 run 数: {miss}")

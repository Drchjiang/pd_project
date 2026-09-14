# -*- coding: utf-8 -*-
import re
from docx import Document
en = Document(r"D:\内推\社招\chjiang\resume_template\数字后端工程师简历模板(英文版).docx")
cn_re = re.compile(r"[一-鿿]")
found = []
for p in en.paragraphs:
    if cn_re.search(p.text):
        found.append(p.text)
for t in en.tables:
    for row in t.rows:
        for c in row.cells:
            if cn_re.search(c.text):
                found.append(c.text)
print("英文版残留中文段落数:", len(found))
for f in found:
    print("  ", f)

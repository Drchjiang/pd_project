# -*- coding: utf-8 -*-
"""生成：芯片数字后端工程师（3年社招）简历模板 .docx"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_TAB_ALIGNMENT, WD_LINE_SPACING, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = r"D:\内推\社招\chjiang\resume_template\数字后端工程师简历模板.docx"

CN = "宋体"        # 正文中文字体
CN_H = "黑体"      # 标题中文字体
EN = "Arial"
BODY_PT = 10.5
DARK = RGBColor(0x22, 0x22, 0x22)
GRAY = RGBColor(0x55, 0x55, 0x55)
LINE_PT = 15       # 固定行距

doc = Document()

# ---------- 页面 ----------
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.9), Cm(1.9)
sec.left_margin, sec.right_margin = Cm(2.2), Cm(2.2)
CONTENT_W = sec.page_width - sec.left_margin - sec.right_margin  # 16.6cm

# ---------- Normal 样式 ----------
normal = doc.styles["Normal"]
normal.font.name = EN
normal.font.size = Pt(BODY_PT)
normal.font.color.rgb = DARK
normal.element.rPr.rFonts.set(qn("w:eastAsia"), CN)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(2)
normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
normal.paragraph_format.line_spacing = Pt(LINE_PT)


def style_run(run, cn=CN, en=EN, size=BODY_PT, bold=False, italic=False, color=DARK):
    run.font.name = en
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), cn)


def para(space_after=2, space_before=0, align=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(LINE_PT)
    if align is not None:
        p.alignment = align
    return p


def add_bottom_border(p, color="666666", sz="6"):
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def section_title(text):
    p = para(space_after=3, space_before=7)
    r = p.add_run(text)
    style_run(r, cn=CN_H, en=EN, size=12.5, bold=True)
    add_bottom_border(p)
    return p


def lr_line(left_text, right_text, left_bold=False, left_size=BODY_PT,
            right_size=BODY_PT, color=DARK):
    """左文字 + 右对齐时间（制表位）"""
    p = para()
    p.paragraph_format.tab_stops.add_tab_stop(CONTENT_W, WD_TAB_ALIGNMENT.RIGHT)
    r1 = p.add_run(left_text)
    style_run(r1, bold=left_bold, size=left_size, color=color)
    r2 = p.add_run("\t" + right_text)
    style_run(r2, size=right_size, color=color)
    return p


def bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    pf = p.paragraph_format
    pf.left_indent = Cm(0.62)
    pf.first_line_indent = Cm(-0.32)
    pf.space_before = Pt(0)
    pf.space_after = Pt(1.5)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(LINE_PT)
    if bold_prefix:
        r0 = p.add_run(bold_prefix)
        style_run(r0, bold=True)
    r = p.add_run(text)
    style_run(r)
    return p


def hint(text):
    p = para(space_after=1)
    r = p.add_run(text)
    style_run(r, size=9, italic=True, color=GRAY)
    return p

# ========== 1. 姓名 + 个人信息 ==========
p = para(space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
r = p.add_run("XX  X X")   # 姓名
style_run(r, cn=CN_H, size=21, bold=True)

p = para(space_after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
info = "电话：XXXXX    邮箱：XX@XX.com    工作年限：3年    现居城市：XX市    求职状态：在职-暂不考虑/离职-随时到岗"
r = p.add_run(info)
style_run(r, size=10)

# ========== 2. 求职意向（无边框表格） ==========
section_title("求职意向")
t = doc.add_table(rows=2, cols=4)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = True
items = [
    ("应聘职位：", "数字后端工程师", "期望行业：", "芯片/半导体"),
    ("期望城市：", "XX市", "期望薪资：", "XX-XXK/月"),
]
for ri, row in enumerate(t.rows):
    for ci, cell in enumerate(row.cells):
        cp = cell.paragraphs[0]
        cp.paragraph_format.space_after = Pt(1)
        cp.paragraph_format.line_spacing = Pt(14)
        rr = cp.add_run(items[ri][ci])
        style_run(rr, size=10, bold=(ci % 2 == 0))
# 到岗时间单独一行
p = para(space_after=2, space_before=1)
r = p.add_run("到岗时间：XX周内到岗        工作性质：全职")
style_run(r, size=10)

# ========== 3. 专业技能 ==========
section_title("专业技能")
skills = [
    ("设计流程：", "熟悉数字芯片后端全流程（Netlist→GDSII），包括逻辑综合（Design Compiler/Genus）、Floorplan、电源网络规划、布局布线（ICC2/Innovus）、时钟树综合（CTS）、功能/时序 ECO；"),
    ("时序收敛：", "熟练使用 PrimeTime/Tempus 进行静态时序分析（STA），熟悉 MMMC 多模式多工艺角签核，能独立完成 setup/hold 违例分析与修复；"),
    ("时钟树：", "能独立完成 CTS 与时钟树优化，熟悉 useful skew、clock gating，可平衡 skew、功耗与时序；"),
    ("物理验证：", "熟练使用 Calibre 完成 DRC/LVS/ERC/Antenna 检查与违例修复，保证 signoff 干净；"),
    ("低功耗设计：", "熟悉 UPF 低功耗流程、多电压域（MV）设计，了解隔离单元、电平转换单元、电源开关单元插入及功耗分析；"),
    ("可靠性签核：", "了解 IR-Drop/EM 分析（Redhawk/Voltus），能协同修复电压降与电迁移违例；"),
    ("验证与DFT：", "了解形式验证（Formality/Conformal）；能配合 DFT 完成扫描链插入、ATPG 及测试友好的物理实现；"),
    ("脚本能力：", "熟练使用 Tcl/Python/Perl/Shell 编写流程自动化与结果检查脚本，熟悉 Linux 环境与 Git 版本管理；"),
    ("工艺经验：", "具备 XXnm/XXnm 工艺节点后端实现经验，参与 X 颗芯片流片，其中 X 颗一次流片成功。【按实际删改】"),
]
for pre, txt in skills:
    bullet(txt, bold_prefix=pre)

# ========== 4. 工作经历 ==========
section_title("工作经历")
lr_line("XX半导体（XX）有限公司", "XX年XX月 - XX年XX月", left_bold=True)
lr_line("数字后端工程师  |  XX部门", "XX市", color=GRAY)
p = para(space_after=1)
r0 = p.add_run("公司简介：")
style_run(r0, bold=True, size=10)
r = p.add_run("公司为专注于 XX（如 SoC/通信基带/AI 加速/MCU）芯片的 Fabless 设计企业，核心产品为 XX，团队约 XX 人。【一句话即可，可删】")
style_run(r, size=10, color=GRAY)
work_duty = [
    "负责 XX 模块从 Netlist 到 GDSII 的后端物理实现，完成 Floorplan、电源网络、PnR、CTS 与 ECO，模块规模约 XX 万门、目标频率 XX MHz；",
    "建立并维护模块 MMMC 时序约束与 STA 环境，使用 PrimeTime 定位关键路径，通过 cell sizing、buffer 插入、布局/布线优化等手段完成 setup/hold 收敛，将 WNS 由 -XX ps 修复至 +XX ps；",
    "独立完成时钟树综合与优化，将 clock skew 控制在 XX ps 内、insertion delay 约 XX ps，并通过时钟门控降低时钟功耗 XX%；",
    "使用 Calibre 完成 DRC/LVS/Antenna 检查，分析并修复 XX 条违例，最终实现零违例 signoff；",
    "基于 UPF 完成多电压域实现及隔离/电平转换/电源开关单元插入，模块动态、静态功耗合计降低 XX%；",
    "使用 Tcl/Python 开发 XX 个自动化脚本，覆盖网表处理、数据准备与结果检查，后端交付效率提升 XX%；",
    "与前端、DFT、模拟及封装团队对接接口、约束与时序预算，参与设计评审，保障模块按计划节点交付。",
]
for d in work_duty:
    bullet(d)

hint("【第二段工作经历模板，如只有一段请整体删除】")
lr_line("XX科技有限公司", "XX年XX月 - XX年XX月", left_bold=True)
lr_line("助理/初级数字后端工程师", "XX市", color=GRAY)
for d in [
    "协助完成 XX 模块布局布线、CTS 辅助优化及物理验证违例修复；",
    "负责流程数据准备、脚本维护与结果汇总，参与 X 次流片交付。",
]:
    bullet(d)

# ========== 5. 项目经历 ==========
section_title("项目经历")

# 项目一
lr_line("项目一：XXnm XX 芯片（SoC/通信基带）后端设计", "XX年XX月 - XX年XX月", left_bold=True)
lr_line("角色：核心模块后端负责人    规模：XX 万门 / 主频 XX MHz", "工艺：XXnm", color=GRAY)
p = para(space_after=1)
r0 = p.add_run("项目背景：")
style_run(r0, bold=True, size=10)
r = p.add_run("该芯片面向 XX 应用，集成 XX、XX 等模块，采用 XXnm 工艺，后端周期仅 XX 个月，面临高频、低功耗与布线拥塞等多重挑战。")
style_run(r, size=10)
p = para(space_after=1)
r = p.add_run("主要职责：")
style_run(r, bold=True, size=10)
for d in [
    "负责 XX 核心模块后端全流程，制定 Floorplan 与电源网络方案，完成 PnR、CTS 与 ECO；",
    "主导模块时序收敛，分析多 corner 关键路径，分层修复 setup/hold 违例共 XX 条；",
    "协同顶层完成模块集成与时序预算划分，解决接口处布线拥塞与物理验证问题；",
    "编写自动化检查脚本与交付 checklist，规范模块交付质量。",
]:
    bullet(d)
p = para(space_after=1)
r = p.add_run("项目成果：")
style_run(r, bold=True, size=10)
for d in [
    "模块频率达到 XX MHz，WNS/TNS 全部收敛，功耗较预算降低 XX%；",
    "零 DRC/LVS 违例交付，支撑芯片一次流片成功且功能测试通过；",
    "沉淀的流程与 checklist 被后续 X 个项目复用，节省约 XX 人日。",
]:
    bullet(d)

# 项目二
lr_line("项目二：XX 模块时序收敛与低功耗/可靠性优化专项", "XX年XX月 - XX年XX月", left_bold=True)
lr_line("角色：专项优化负责人    规模：XX 万门", "工艺：XXnm", color=GRAY)
p = para(space_after=1)
r0 = p.add_run("项目背景：")
style_run(r0, bold=True, size=10)
r = p.add_run("针对 XX 芯片 XX 模块在 XXnm 工艺下 WNS 违例严重、局部 IR-Drop 超标的问题，开展专项优化以保障流片节点。")
style_run(r, size=10)
p = para(space_after=1)
r = p.add_run("主要职责：")
style_run(r, bold=True, size=10)
for d in [
    "重新梳理关键时序路径，结合 useful skew、cell sizing、布局与逻辑重组进行增量优化；",
    "优化时钟树结构与时钟门控，降低动态功耗；",
    "使用 Redhawk/Voltus 分析 IR-Drop/EM，通过补加电源 stripe、插入 decap 单元修复电压热点。",
]:
    bullet(d)
p = para(space_after=1)
r = p.add_run("项目成果：")
style_run(r, bold=True, size=10)
for d in [
    "WNS 由 -XX ps 提升至 +XX ps，hold 违例清零；",
    "最差 IR-Drop 由 XX% 降至 XX% 以内，模块功耗降低 XX%；",
    "专项仅用 XX 周完成，保障芯片按时流片。",
]:
    bullet(d)

# ========== 6. 教育经历 ==========
section_title("教育经历")
lr_line("XX大学（985/211/双一流）", "XX年XX月 - XX年XX月", left_bold=True)
lr_line("XX专业  硕士研究生/本科", "XX市", color=GRAY)
bullet("GPA：XX/XX（专业前 XX%）；主修课程：数字集成电路设计、半导体物理、Verilog HDL、低功耗设计、CMOS 模拟电路等；【按实际填写】")
bullet("在校荣誉：XX 奖学金、XX 竞赛 XX 奖。【无可删】")

# ========== 7. 证书 & 自我评价 ==========
section_title("荣誉证书 / 自我评价")
bullet("证书：XX EDA 工具认证、英语 CET-X（XX 分）、XX。【无可删】")
p = para(space_after=2)
r0 = p.add_run("自我评价：")
style_run(r0, bold=True)
r = p.add_run("拥有 3 年数字后端设计经验，熟悉 Netlist→GDSII 全流程及主流 EDA 工具，具备 XXnm 工艺 X 颗芯片流片经验；在时序收敛、时钟树、物理验证与低功耗设计上可独立承担模块级交付；善用脚本提升流程效率，工作严谨、沟通协作顺畅，可快速上手新项目。")
style_run(r)

doc.save(OUT)
print("saved:", OUT)
